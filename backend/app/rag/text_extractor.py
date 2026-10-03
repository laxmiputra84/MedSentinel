import os
import csv
import re
from typing import Dict, Any, List, Optional
import pypdf
import docx
import openpyxl

class TextExtractor:
    @staticmethod
    def clean_text(raw_text: str) -> str:
        """Cleans and standardizes extracted text content."""
        if not raw_text:
            return ""
        # Replace non-breaking spaces and tabs
        text = raw_text.replace("\xa0", " ").replace("\t", " ")
        # Remove multiple consecutive blank lines
        text = re.sub(r"\n\s*\n", "\n\n", text)
        # Remove multiple consecutive spaces
        text = re.sub(r"[ ]{2,}", " ", text)
        return text.strip()

    @staticmethod
    def extract_from_pdf(file_path: str) -> str:
        """Extracts text content from PDF documents."""
        text_parts = []
        try:
            reader = pypdf.PdfReader(file_path)
            for i, page in enumerate(reader.pages):
                page_text = page.extract_text()
                if page_text:
                    text_parts.append(f"--- Page {i + 1} ---\n{page_text}")
        except Exception as e:
            print(f"Error reading PDF {file_path}: {e}")
        return TextExtractor.clean_text("\n\n".join(text_parts))

    @staticmethod
    def extract_from_docx(file_path: str) -> str:
        """Extracts text content from Word DOCX documents."""
        text_parts = []
        try:
            doc = docx.Document(file_path)
            for p in doc.paragraphs:
                if p.text.strip():
                    text_parts.append(p.text.strip())
            # Extract tables as well
            for table in doc.tables:
                for row in table.rows:
                    row_data = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                    if row_data:
                        text_parts.append(" | ".join(row_data))
        except Exception as e:
            print(f"Error reading DOCX {file_path}: {e}")
        return TextExtractor.clean_text("\n\n".join(text_parts))

    @staticmethod
    def extract_from_csv(file_path: str) -> str:
        """Extracts structured rows from CSV documents."""
        lines = []
        try:
            with open(file_path, mode="r", encoding="utf-8", errors="replace") as f:
                reader = csv.reader(f)
                headers = next(reader, None)
                if headers:
                    lines.append("Columns: " + ", ".join(headers))
                for row in reader:
                    if row:
                        if headers and len(headers) == len(row):
                            row_str = "; ".join([f"{h}: {v}" for h, v in zip(headers, row) if v.strip()])
                            lines.append(row_str)
                        else:
                            lines.append(" | ".join(row))
        except Exception as e:
            print(f"Error reading CSV {file_path}: {e}")
        return TextExtractor.clean_text("\n".join(lines))

    @staticmethod
    def extract_from_xlsx(file_path: str) -> str:
        """Extracts structured data from Excel XLSX workbooks."""
        lines = []
        try:
            wb = openpyxl.load_workbook(file_path, read_only=True, data_only=True)
            for sheet in wb.sheetnames:
                ws = wb[sheet]
                lines.append(f"=== Sheet: {sheet} ===")
                for row in ws.iter_rows(values_only=True):
                    row_vals = [str(v).strip() for v in row if v is not None and str(v).strip()]
                    if row_vals:
                        lines.append(" | ".join(row_vals))
        except Exception as e:
            print(f"Error reading XLSX {file_path}: {e}")
        return TextExtractor.clean_text("\n".join(lines))

    @staticmethod
    def extract_from_txt(file_path: str) -> str:
        """Extracts plain text files."""
        try:
            with open(file_path, "r", encoding="utf-8", errors="replace") as f:
                return TextExtractor.clean_text(f.read())
        except Exception as e:
            print(f"Error reading TXT {file_path}: {e}")
            return ""

    @staticmethod
    def extract_text(file_path: str, file_type: Optional[str] = None) -> str:
        """Detects format and extracts cleaned text from document."""
        if not os.path.exists(file_path):
            return ""

        ext = (file_type or os.path.splitext(file_path)[1]).lower().replace(".", "")

        if ext == "pdf":
            return TextExtractor.extract_from_pdf(file_path)
        elif ext in ["docx", "doc"]:
            return TextExtractor.extract_from_docx(file_path)
        elif ext == "csv":
            return TextExtractor.extract_from_csv(file_path)
        elif ext in ["xlsx", "xls"]:
            return TextExtractor.extract_from_xlsx(file_path)
        else:
            return TextExtractor.extract_from_txt(file_path)
